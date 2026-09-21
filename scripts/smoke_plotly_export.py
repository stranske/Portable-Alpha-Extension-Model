"""Run the real CLI chart export and require an image in the output workbook.

Run with an environment containing Plotly, Kaleido and a compatible Chrome:
    python scripts/smoke_plotly_export.py

The CLI must run without its CI/pytest chart bypass. An exit-zero workbook with
no Summary image is a failure, even if the chart renderer swallowed an exception.
"""

from __future__ import annotations

import importlib.metadata
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl


def main() -> None:
    repo = Path(__file__).resolve().parents[1]
    env = os.environ.copy()
    env.pop("CI", None)
    env.pop("PYTEST_CURRENT_TEST", None)
    with tempfile.TemporaryDirectory(prefix="paem-plotly-smoke-") as output_dir:
        workbook = Path(output_dir) / "sweep.xlsx"
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "pa_core.cli",
                "--config",
                "examples/scenarios/my_first_scenario.yml",
                "--index",
                "data/sp500tr_fred_divyield.csv",
                "--mode",
                "capital",
                "--output",
                str(workbook),
                "--seed",
                "42",
            ],
            cwd=repo,
            env=env,
            capture_output=True,
            text=True,
            timeout=180,
            check=True,
        )
        summary = openpyxl.load_workbook(workbook)["Summary"]
        if summary.max_row <= 1 or len(summary._images) < 1:
            raise RuntimeError(
                "CLI produced a workbook without a populated Summary and embedded chart; "
                f"stdout={result.stdout!r} stderr={result.stderr!r}"
            )
        print(
            f"PASS plotly={importlib.metadata.version('plotly')} "
            f"summary_rows={summary.max_row} images={len(summary._images)}"
        )


if __name__ == "__main__":
    main()
